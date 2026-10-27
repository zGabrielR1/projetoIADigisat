"""Evidence exports from stored results or an explicitly requested fresh full query."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
import csv
import json
import os
import re
from xml.sax.saxutils import escape

from fastapi import HTTPException
from pydantic import BaseModel, Field
import insight_store as store
from insight_plans import Plan
from insight_engine import all_results

_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="evidence-export")
LABELS = {"Data": "Emissão", "Numero": "Número", "Serie": "Série", "Empresa": "Empresa", "Tipo": "Tipo",
          "Situacao": "Situação", "Itens": "Itens do documento", "Vendedor": "Itens do vendedor", "Outros": "Outros/sem vendedor",
          "Pagamentos": "Pagamentos do documento", "Diferenca": "Itens - pagamentos", "Observacao": "Conferência",
          "Produto": "Produto", "Quantidade": "Quantidade", "Valor": "Valor", "Documento": "Documento", "total": "Total",
          "items": "Itens", "documents": "Documentos", "_id": "Grupo", "ticket": "Ticket médio", "participacao": "Participação"}


class ExportRequest(BaseModel):
    run_id: str = Field(min_length=10, max_length=100)
    format: str = Field(default="xlsx", pattern="^(xlsx|pdf|csv)$")
    complete: bool = False


def safe_text(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if value is None:
        return ""
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
        return value[:32767]
    return value


def put_job(key, body):
    store.write("INSERT OR REPLACE INTO jobs VALUES(?,?,?)", (key, store.dump(body), store.now()))


def get_job(key):
    rows = store.read("SELECT body FROM jobs WHERE id=?", (key,))
    return json.loads(rows[0]["body"]) if rows else None


def mutate_job(key, changes=None, *, check_cancel=False):
    """Serialize progress/cancellation updates so one cannot erase the other."""
    with store.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        row = db.execute("SELECT body FROM jobs WHERE id=?", (key,)).fetchone()
        if not row:
            raise HTTPException(404, "Exportação não encontrada.")
        job = json.loads(row["body"])
        if check_cancel and job.get("cancel_requested"):
            raise HTTPException(499, "Exportação cancelada. Nenhum arquivo parcial foi entregue.")
        job.update(changes or {})
        db.execute("UPDATE jobs SET body=?,updated=? WHERE id=?", (store.dump(job), store.now(), key))
        return job


def cancel_job(key):
    with store.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        row = db.execute("SELECT body FROM jobs WHERE id=?", (key,)).fetchone()
        if not row:
            raise HTTPException(404, "Exportação não encontrada.")
        job = json.loads(row["body"])
        if job["status"] in {"queued", "running"}:
            job["cancel_requested"] = True
            if job["status"] == "queued":
                job.update(status="cancelled", error="Exportação cancelada antes da execução.")
            db.execute("UPDATE jobs SET body=?,updated=? WHERE id=?", (store.dump(job), store.now(), key))
        return job


def start_export(req: ExportRequest) -> dict:
    run = store.get_run(req.run_id)
    if not run:
        raise HTTPException(404, "Resultado não encontrado. Execute a consulta novamente.")
    if req.complete and run['plan'].get('module') == 'stock_expiry':
        raise HTTPException(422, 'Para lotes, exporte a página exibida. Uma consulta completa não está disponível neste módulo.')
    key = store.identifier()
    job = {"id": key, "status": "queued", "format": req.format, "records": 0,
           "source": "live_complete" if req.complete else "saved_snapshot", "run_id": req.run_id, "owner_pid": os.getpid()}
    put_job(key, job)
    _pool.submit(build_export, key, req, run)
    return job


def datasets(req, run, key):
    mutate_job(key, check_cancel=True)
    result = run["result"]
    if not req.complete:
        yield "Base", result
        if result.get("comparison"):
            yield "Comparação", result["comparison"]["result"]
        return
    plan = Plan(**run["plan"])
    pages = all_results(plan)
    iterator = iter(pages)
    while True:
        mutate_job(key, check_cancel=True)
        try:
            page = next(iterator)
        except StopIteration:
            break
        yield "Base", page
    if plan.comparison:
        other = plan.model_dump(mode="json")
        other.update(plan.comparison.model_dump(mode="json"))
        other.update(comparison=None, kind="investigation", offset=0)
        other.pop("label", None)
        iterator = iter(all_results(Plan(**other)))
        while True:
            mutate_job(key, check_cancel=True)
            try:
                page = next(iterator)
            except StopIteration:
                break
            yield "Comparação", page


def build_export(key, req, run):
    folder = store.data_dir() / "exports"
    path = folder / f"{key}.{req.format}"
    try:
        mutate_job(key, {"status": "running", "started_at": store.now()}, check_cancel=True)
        folder.mkdir(exist_ok=True)
        if req.format == "xlsx":
            write_xlsx(path, req, run, key)
        elif req.format == "pdf":
            write_pdf(path, req, run, key)
        else:
            write_csv(path, req, run, key)
        job = mutate_job(key, {"status": "succeeded", "finished_at": store.now(), "filename": path.name,
                   "download": f"/api/workspace/exports/{key}/download"}, check_cancel=True)
        store.event("export", f"format={req.format}; complete={req.complete}; records={job['records']}")
    except Exception as exc:
        message = str(exc.detail) if isinstance(exc, HTTPException) else f"Falha ao gerar evidências: {type(exc).__name__}. Tente novamente ou reduza o período."
        mutate_job(key, {"status": "cancelled" if isinstance(exc, HTTPException) and exc.status_code == 499 else "failed", "error": message})
        if path.exists():
            path.unlink()


def update_progress(key, count):
    mutate_job(key, {"records": count}, check_cancel=True)


def metadata(run, req):
    result = run["result"]
    rows = [["DigiSat Insight", "Evidências da consulta"], ["Resultado original", run["created"]],
            ["Pergunta original", result.get("question", "Consulta pelos filtros")],
            ["Exportação", store.now()], ["Origem", "Nova consulta completa à base atual" if req.complete else "Cópia do resultado salvo (não recalculado)"],
            ["Consistência", "Leituras ao longo da exportação; não é snapshot transacional do MongoDB." if req.complete else "Valores preservados no momento da consulta original."],
            ["Aplicação", result.get("provenance", {}).get("version", "")],
            ["Banco", result.get("provenance", {}).get("database", "")]]
    rows.extend([["Plano", text] for text in result.get("interpretation", {}).get("summary", [])])
    rows.extend([["Conclusão original" if req.complete else "Conclusão", text] for text in result.get("investigation", {}).get("findings", [])])
    rows.extend([["Aviso do resultado original" if req.complete else "Aviso", text] for text in result.get("warnings", [])])
    return rows


def write_xlsx(path, req, run, key):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook(write_only=True)
    summary = book.create_sheet("Resumo e método")
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 110
    for a, b in metadata(run, req):
        cell = WriteOnlyCell(summary, value=safe_text(b))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        summary.append([safe_text(a), cell])
    sheets = {}
    item_counts = {}
    count = 0
    for label, data in datasets(req, run, key):
        rows = data.get("rows", [])
        columns = list(dict.fromkeys(k for row in rows for k in row))
        if label not in sheets:
            if req.complete:
                for finding in data.get("investigation", {}).get("findings", []):
                    summary.append(["Conclusão atual - " + label, safe_text(finding)])
            sheet = book.create_sheet(label + (" - lotes" if run['result'].get('module') == 'stock_expiry' else " - vendas"))
            sheet.freeze_panes = "A2"
            for index in range(len(columns)):
                from openpyxl.utils import get_column_letter
                sheet.column_dimensions[get_column_letter(index + 1)].width = 22
            headers = []
            for column in columns:
                cell = WriteOnlyCell(sheet, value=LABELS.get(column, column))
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="16624B")
                headers.append(cell)
            sheet.append(headers)
            items = book.create_sheet(label + " - itens")
            items.freeze_panes = "A2"
            for col, width in [("A", 20), ("B", 12), ("C", 65), ("D", 18), ("E", 35), ("F", 20)]:
                items.column_dimensions[col].width = width
            items.append(["Documento", "Item", "Produto", "Quantidade", "Vendedor", "Valor"])
            sheets[label] = (sheet, items, columns, 0)
            item_counts[label] = 0
        sheet, items, columns, written = sheets[label]
        if written + len(rows) >= 1048575:
            raise HTTPException(422, "O resultado excede o limite de uma planilha Excel. Divida a consulta.")
        for row in rows:
            cells = []
            for column in columns:
                value = safe_text(row.get(column))
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, (float, int)) and column in {"Itens", "Vendedor", "Outros", "Pagamentos", "Diferenca", "Valor", "total", "ticket"}:
                    cell.number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
                cells.append(cell)
            sheet.append(cells)
        for doc in data.get("investigation", {}).get("details", []):
            item_counts[label] += len(doc["items"])
            if item_counts[label] >= 1048575:
                raise HTTPException(422, "Os itens excedem o limite de uma planilha Excel. Divida a consulta ou exporte CSV.")
            for item in doc["items"]:
                items.append([safe_text(doc["document"]), safe_text(item.get("Numero")), safe_text(item.get("Produto")),
                              safe_text(item.get("Quantidade")), safe_text(item.get("Vendedor")), safe_text(item.get("Valor"))])
        sheets[label] = (sheet, items, columns, written + len(rows))
        count += len(rows)
        update_progress(key, count)
    book.save(path)


def write_csv(path, req, run, key):
    count = 0
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output, delimiter=";")
        writer.writerows([[safe_text(value) for value in row] for row in metadata(run, req)])
        for label, data in datasets(req, run, key):
            rows = data.get("rows", [])
            columns = list(dict.fromkeys(k for row in rows for k in row))
            writer.writerow([])
            writer.writerow([label])
            if req.complete and (data.get("pagination") or {}).get("offset", 0) == 0:
                for finding in data.get("investigation", {}).get("findings", []):
                    writer.writerow(["Conclusão atual", safe_text(finding)])
            writer.writerow([LABELS.get(k, k) for k in columns])
            for row in rows:
                writer.writerow([safe_text(row.get(k)) for k in columns])
            count += len(rows)
            update_progress(key, count)


def write_pdf(path, req, run, key):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    font = "Helvetica"
    font_path = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "segoeui.ttf"
    if font_path.exists():
        if "InsightSegoe" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("InsightSegoe", str(font_path)))
        font = "InsightSegoe"
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
    body = ParagraphStyle("EvidenceBody", fontName=font, fontSize=8, leading=11, spaceAfter=5)
    small = ParagraphStyle("EvidenceCell", fontName=font, fontSize=7, leading=9)
    def p(value, style=body):
        text = str(value if value is not None else "").translate(str.maketrans({"—": "-", "–": "-", "−": "-", "‑": "-"}))
        return Paragraph(escape(text).replace("\n", "<br/>"), style)
    story = [p("DigiSat Insight — Evidências da consulta", styles["Title"]), Spacer(1, 12)]
    for a, b in metadata(run, req)[1:]:
        story.append(p(f"{a}: {b}"))
    count = 0
    item_count = 0
    for label, data in datasets(req, run, key):
        rows = data.get("rows", [])
        if count + len(rows) > 2000:
            raise HTTPException(422, "O PDF excederia 2.000 vendas. Use a exportação Excel/CSV completa ou reduza o período; nenhum PDF parcial foi entregue.")
        story.extend([Spacer(1, 12), p(label + " — documentos", styles["Heading2"])])
        if req.complete and (data.get("pagination") or {}).get("offset", 0) == 0:
            for finding in data.get("investigation", {}).get("findings", []):
                story.append(p("Conclusão atual: " + finding))
        available = list(dict.fromkeys(k for row in rows for k in row))
        preferred = [k for k in ("Serie", "Numero", "Data", "Itens", "Vendedor", "Outros", "Pagamentos", "Diferenca") if k in available]
        columns = preferred or available[:7]
        def fmt(k, value):
            if k == "Data" and isinstance(value, str):
                try:
                    return datetime.fromisoformat(value.replace("Z", "+00:00")).strftime("%d/%m/%Y %H:%M")
                except ValueError:
                    pass
            if isinstance(value, (int, float)) and k in {"Itens", "Vendedor", "Outros", "Pagamentos", "Diferenca", "Valor", "total", "ticket"}:
                import app
                return app.brl(value)
            return safe_text(value)
        if columns:
            values = [[p(LABELS.get(k, k), small) for k in columns]]
            values += [[p(fmt(k, row.get(k)), small) for k in columns] for row in rows]
            table = Table(values, colWidths=[750 / len(columns)] * len(columns), repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2eee7")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#ccd8d0")),
                                       ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.append(table)
        for doc in data.get("investigation", {}).get("details", []):
            item_count += len(doc["items"])
            if item_count > 20000:
                raise HTTPException(422, "O PDF excederia 20.000 itens. Use Excel/CSV ou reduza o período.")
            story.extend([Spacer(1, 9), p("Venda " + doc["document"], styles["Heading3"])])
            values = [[p(k, small) for k in ["Item", "Produto", "Quantidade", "Vendedor", "Valor"]]]
            values += [[p(item.get("Numero"), small), p(item.get("Produto"), small), p(item.get("Quantidade"), small),
                        p(item.get("Vendedor"), small), p(fmt("Valor", item.get("Valor")), small)] for item in doc["items"]]
            table = Table(values, colWidths=[40, 330, 70, 210, 100], repeatRows=1)
            table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff4ef")),
                                       ("LINEBELOW", (0, 0), (-1, -1), .2, colors.lightgrey)]))
            story.append(table)
        count += len(rows)
        update_progress(key, count)
    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawString(36, 20, "DigiSat Insight | Consulta somente leitura | " + store.now()[:10])
        canvas.drawRightString(805, 20, f"Página {doc.page}")
    document = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=36, rightMargin=36, topMargin=30, bottomMargin=35)
    document.build(story, onFirstPage=footer, onLaterPages=footer)
